"""Excel exports, printable PDFs and e-mailing documents."""

import io
import os
import re
from datetime import date, datetime
from decimal import Decimal

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, EmailStr, Field
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image as RLImage
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import or_
from sqlalchemy.orm import Session

from ..database import get_db
from ..deps import get_company, get_current_user
from ..models import Company, QBRecord, QBReport, User
from ..qb import builder
from ..qb.registry import BY_KEY, REPORTS, Entity
from ..services import audit
from ..services.notify import send_email_now

router = APIRouter(prefix="/api/files", tags=["Export, print, e-mail"])
LOGO = os.path.join(os.path.dirname(os.path.dirname(__file__)), "assets", "logo.png")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _entity(key: str, user: User) -> Entity:
    ent = BY_KEY.get(key)
    if not ent:
        raise HTTPException(404, "Unknown entity")
    if f"{ent.module}.view" not in (user.role.permissions or []):
        raise HTTPException(403, f"Missing permission: {ent.module}.view")
    return ent


def _like(q: str) -> str:
    return "%" + q.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_") + "%"


def _safe(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "_", name).strip("_") or "export"


def _xlsx(title: str, header: list[str], rows: list[list], money_cols: set[int]) -> StreamingResponse:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0A5C80")
    for r in rows:
        ws.append(r)
    for i, h in enumerate(header, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1] if i - 1 < len(r) and r[i - 1] is not None else "")) for r in rows[:500]])
        ws.column_dimensions[get_column_letter(i)].width = min(60, width + 2)
        if i - 1 in money_cols:
            for cell in ws[get_column_letter(i)][1:]:
                cell.number_format = "#,##0.00"
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"{_safe(title)}_{date.today():%Y%m%d}.xlsx"
    return StreamingResponse(buf, media_type=XLSX, headers={"Content-Disposition": f"attachment; filename={fname}"})


def _num(v):
    try:
        return float(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return v


# ---------------------------------------------------------------- Excel

@router.get("/export/{entity}")
def export_list(entity: str, q: str | None = None, active: bool | None = True, date_from: date | None = None,
                date_to: date | None = None, db: Session = Depends(get_db), user: User = Depends(get_current_user),
                co: Company = Depends(get_company)):
    ent = _entity(entity, user)
    query = db.query(QBRecord).filter(QBRecord.company_id == co.company_id, QBRecord.entity == ent.key,
                                      QBRecord.deleted.is_(False))
    if q:
        for word in q.split():
            query = query.filter(or_(QBRecord.name.ilike(_like(word)), QBRecord.party_name.ilike(_like(word))))
    if ent.kind == "list" and active is not None:
        query = query.filter(QBRecord.is_active.is_(active))
    if date_from:
        query = query.filter(QBRecord.txn_date >= date_from)
    if date_to:
        query = query.filter(QBRecord.txn_date <= date_to)
    order = (QBRecord.txn_date.desc().nullslast(),) if ent.kind == "txn" else (QBRecord.name,)
    fields = [f for f in ent.fields if f.type not in ("address",)]
    header = (["No.", "Date", "Name", "Amount"] if ent.kind == "txn" else ["Name"]) + [f.label for f in fields]
    rows = []
    for r in query.order_by(*order).limit(50000).all():
        data = builder.normalize(ent, r.data or {})
        base = [r.name, r.txn_date, r.party_name, float(r.amount) if r.amount is not None else None] \
            if ent.kind == "txn" else [r.name]
        vals = []
        for f in fields:
            v = builder.get(data, f.path)
            if isinstance(v, dict):
                v = v.get("FullName") or v.get("ListID")
            vals.append(_num(v) if f.type in ("money", "decimal", "int") else v)
        rows.append(base + vals)
    money = {3} if ent.kind == "txn" else set()
    money |= {len(header) - len(fields) + i for i, f in enumerate(fields) if f.type == "money"}
    audit.log(db, user.user_id, "export", ent.key, "excel", None, {"rows": len(rows), "company": co.name})
    db.commit()
    return _xlsx(f"{co.name} {ent.plural}", header, rows, money)


@router.get("/report/{report_id}")
def export_report(report_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user),
                  co: Company = Depends(get_company)):
    if "reports.view" not in (user.role.permissions or []):
        raise HTTPException(403, "Missing permission: reports.view")
    r = db.get(QBReport, report_id)
    if not r or r.company_id != co.company_id or not r.result:
        raise HTTPException(404, "Report not ready")
    res = r.result
    rows = [[c if i == 0 else _num(c) for i, c in enumerate(row["cells"])] for row in res["rows"]]
    label = REPORTS.get(r.report_type, ("", r.report_type))[1]
    return _xlsx(f"{co.name} {label}", res["columns"] or [""], rows, set(range(1, len(res["columns"]))))


# ---------------------------------------------------------------- PDF documents

TITLES = {"invoice": "INVOICE", "estimate": "ESTIMATE", "sales_order": "SALES ORDER", "sales_receipt": "SALES RECEIPT",
          "credit_memo": "CREDIT MEMO", "purchase_order": "PURCHASE ORDER", "bill": "BILL", "item_receipt": "ITEM RECEIPT",
          "vendor_credit": "VENDOR CREDIT", "check": "CHECK", "credit_card_charge": "CREDIT CARD CHARGE",
          "credit_card_credit": "CREDIT CARD CREDIT", "receive_payment": "PAYMENT RECEIPT", "deposit": "DEPOSIT",
          "journal_entry": "JOURNAL ENTRY", "statement_charge": "STATEMENT CHARGE"}


def _money(v) -> str:
    try:
        return f"{Decimal(str(v)):,.2f}"
    except Exception:
        return ""


def _addr(a) -> str:
    if not isinstance(a, dict):
        return ""
    parts = [a.get(k) for k in ("Addr1", "Addr2", "Addr3", "Addr4", "Addr5")]
    city = " ".join(p for p in (a.get("City"), a.get("State"), a.get("PostalCode")) if p)
    return "<br/>".join(p for p in parts + [city, a.get("Country")] if p)


def build_pdf(co: Company, ent: Entity, rec: QBRecord) -> bytes:
    data = rec.data or {}
    st = getSampleStyleSheet()
    small = st["BodyText"].clone("small", fontSize=8.5, leading=11)
    head = st["Title"].clone("head", alignment=2, fontSize=20, textColor=colors.HexColor("#0A5C80"))
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
                            title=f"{TITLES.get(ent.key, ent.label.upper())} {rec.name or ''}")
    story = []
    brand = Paragraph(f"<b>{co.name}</b><br/><font size='8' color='grey'>Synex Group</font>", st["Heading2"])
    if os.path.isfile(LOGO):
        brand = Table([[RLImage(LOGO, width=20 * mm, height=20 * mm), brand]], colWidths=[24 * mm, 71 * mm])
        brand.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    top = Table([[brand, Paragraph(TITLES.get(ent.key, ent.label.upper()), head)]], colWidths=[95 * mm, 79 * mm])
    top.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [top, Spacer(1, 6)]

    party = builder.get(data, ent.party_path) if ent.party_path else None
    party_name = party.get("FullName") if isinstance(party, dict) else ""
    address = _addr(data.get("BillAddress") or data.get("VendorAddress") or data.get("Address"))
    meta_rows = [["No.", rec.name or "—"], ["Date", data.get("TxnDate", "")]]
    for k, label in (("DueDate", "Due date"), ("PONumber", "P.O. no."), ("TermsRef", "Terms"), ("ExpectedDate", "Expected")):
        v = data.get(k)
        if v:
            meta_rows.append([label, v.get("FullName") if isinstance(v, dict) else v])
    meta_tbl = Table(meta_rows, colWidths=[28 * mm, 45 * mm])
    meta_tbl.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 9), ("TEXTCOLOR", (0, 0), (0, -1), colors.grey)]))
    left = Paragraph(f"<font color='grey' size='8'>{'BILL TO' if ent.party_path == 'CustomerRef' else 'TO'}</font><br/>"
                     f"<b>{party_name or ''}</b><br/>{address}", st["BodyText"])
    blk = Table([[left, meta_tbl]], colWidths=[100 * mm, 74 * mm])
    blk.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [blk, Spacer(1, 10)]

    # lines (all line types of the entity, in order)
    rows = [["Item / account", "Description", "Qty", "Rate", "Amount"]]
    total = Decimal(0)
    for lt in ent.lines:
        for ln in data.get(lt.ret_tag, []) or []:
            if not isinstance(ln, dict):
                continue
            what = ln.get("ItemRef") or ln.get("AccountRef") or ln.get("ItemInventoryRef") or {}
            qty, rate = ln.get("Quantity", ""), ln.get("Rate") or ln.get("Cost") or ""
            amount = ln.get("Amount") or ln.get("PaymentAmount")
            if amount is None and qty and rate:
                amount = Decimal(str(qty)) * Decimal(str(rate))
            if amount not in (None, ""):
                total += Decimal(str(amount))
            rows.append([Paragraph(what.get("FullName", "") if isinstance(what, dict) else str(what), small),
                         Paragraph(str(ln.get("Desc") or ln.get("Memo") or ""), small),
                         qty, _money(rate) if rate else "", _money(amount) if amount not in (None, "") else ""])
    if len(rows) > 1:
        t = Table(rows, colWidths=[40 * mm, 70 * mm, 16 * mm, 22 * mm, 26 * mm], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6F0F5")), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("ALIGN", (2, 0), (-1, -1), "RIGHT"),
            ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E7EB")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
        story.append(t)

    shown_total = data.get("TotalAmount") or data.get("Subtotal") or data.get("Amount") or data.get("AmountDue") \
        or data.get("CreditAmount") or data.get("DepositTotal") or (total if total else None)
    totals = [["Total", _money(shown_total)]] if shown_total is not None else []
    if data.get("BalanceRemaining") not in (None, ""):
        totals.append(["Balance due", _money(data["BalanceRemaining"])])
    if totals:
        tt = Table(totals, colWidths=[30 * mm, 30 * mm], hAlign="RIGHT")
        tt.setStyle(TableStyle([("ALIGN", (1, 0), (1, -1), "RIGHT"), ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
                                ("FONTSIZE", (0, 0), (-1, -1), 10), ("LINEABOVE", (0, 0), (-1, 0), 0.75, colors.black)]))
        story += [Spacer(1, 8), tt]
    if data.get("Memo"):
        story += [Spacer(1, 14), Paragraph(f"<font color='grey'>Memo:</font> {data['Memo']}", st["BodyText"])]
    story += [Spacer(1, 20), Paragraph(f"<font size='7' color='grey'>Generated by Synex QB Portal from QuickBooks · "
                                       f"{datetime.now():%Y-%m-%d %H:%M}</font>", st["BodyText"])]
    doc.build(story)
    return buf.getvalue()


def _doc_record(db: Session, entity: str, record_id: int, user: User, co: Company) -> tuple[Entity, QBRecord]:
    ent = _entity(entity, user)
    rec = db.get(QBRecord, record_id)
    if not rec or rec.entity != ent.key or rec.company_id != co.company_id:
        raise HTTPException(404, "Record not found")
    return ent, rec


@router.get("/pdf/{entity}/{record_id}")
def pdf(entity: str, record_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user),
        co: Company = Depends(get_company)):
    ent, rec = _doc_record(db, entity, record_id, user, co)
    content = build_pdf(co, ent, rec)
    fname = f"{_safe(ent.label)}_{_safe(rec.name or str(rec.record_id))}.pdf"
    return StreamingResponse(io.BytesIO(content), media_type="application/pdf",
                             headers={"Content-Disposition": f"inline; filename={fname}"})


class EmailIn(BaseModel):
    to: list[EmailStr] = Field(min_length=1, max_length=10)
    subject: str | None = Field(default=None, max_length=200)
    message: str | None = Field(default=None, max_length=5000)


@router.post("/email/{entity}/{record_id}")
def email(entity: str, record_id: int, body: EmailIn, db: Session = Depends(get_db),
          user: User = Depends(get_current_user), co: Company = Depends(get_company)):
    ent, rec = _doc_record(db, entity, record_id, user, co)
    title = f"{TITLES.get(ent.key, ent.label)} {rec.name or ''}".strip()
    subject = body.subject or f"{co.name}: {title.title()}"
    text = body.message or f"Please find attached {title.lower()} from {co.name}."
    try:
        send_email_now([str(t) for t in body.to], subject, text,
                       attachment=(f"{_safe(title)}.pdf", build_pdf(co, ent, rec), "application/pdf"))
    except RuntimeError as e:
        raise HTTPException(400, str(e))
    audit.log(db, user.user_id, "email", rec.record_id, "send", None, {"to": [str(t) for t in body.to], "doc": title})
    db.commit()
    return {"message": f"Sent to {', '.join(str(t) for t in body.to)}"}


# ---------------------------------------------------------------- attachments (stored by the portal)

import os  # noqa: E402
import uuid  # noqa: E402

from fastapi import File, UploadFile  # noqa: E402
from fastapi.responses import FileResponse  # noqa: E402

from ..config import get_settings  # noqa: E402
from ..models import Attachment  # noqa: E402

BLOCKED_EXT = {".exe", ".bat", ".cmd", ".com", ".msi", ".ps1", ".vbs", ".js", ".jar", ".scr", ".dll", ".html", ".htm",
               ".svg"}


class AttachmentOut(BaseModel):
    attachment_id: int
    filename: str
    content_type: str
    size: int
    uploaded_by: str | None
    uploaded_at: datetime


def _att_out(a: Attachment) -> AttachmentOut:
    return AttachmentOut(attachment_id=a.attachment_id, filename=a.filename, content_type=a.content_type, size=a.size,
                         uploaded_by=(a.uploader.full_name or a.uploader.username) if a.uploader else None,
                         uploaded_at=a.uploaded_at)


@router.get("/attachments/{entity}/{record_id}", response_model=list[AttachmentOut])
def list_attachments(entity: str, record_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user),
                     co: Company = Depends(get_company)):
    ent, rec = _doc_record(db, entity, record_id, user, co)
    rows = (db.query(Attachment).filter_by(company_id=co.company_id, entity=ent.key, qb_id=rec.qb_id)
            .order_by(Attachment.uploaded_at.desc()).all())
    return [_att_out(a) for a in rows]


@router.post("/attachments/{entity}/{record_id}", response_model=AttachmentOut, status_code=201)
async def upload_attachment(entity: str, record_id: int, file: UploadFile = File(...), db: Session = Depends(get_db),
                            user: User = Depends(get_current_user), co: Company = Depends(get_company)):
    ent, rec = _doc_record(db, entity, record_id, user, co)
    if f"{ent.module}.edit" not in (user.role.permissions or []) and f"{ent.module}.create" not in (user.role.permissions or []):
        raise HTTPException(403, f"Missing permission: {ent.module}.edit")
    name = os.path.basename(file.filename or "file")[:255]
    if os.path.splitext(name)[1].lower() in BLOCKED_EXT:
        raise HTTPException(422, "This file type is not allowed")
    limit = get_settings().max_attachment_mb * 1024 * 1024
    content = await file.read(limit + 1)
    if len(content) > limit:
        raise HTTPException(413, f"File larger than {get_settings().max_attachment_mb} MB")
    folder = os.path.join(get_settings().attachments_dir, str(co.company_id))
    os.makedirs(folder, exist_ok=True)
    stored = f"{uuid.uuid4().hex}{os.path.splitext(name)[1].lower()}"
    with open(os.path.join(folder, stored), "wb") as fh:
        fh.write(content)
    a = Attachment(company_id=co.company_id, entity=ent.key, qb_id=rec.qb_id, filename=name,
                   content_type=file.content_type or "application/octet-stream", size=len(content),
                   stored_as=os.path.join(str(co.company_id), stored), uploaded_by=user.user_id)
    db.add(a)
    db.flush()
    audit.log(db, user.user_id, "attachment", a.attachment_id, "upload", None, {"file": name, "record": rec.name})
    db.commit()
    return _att_out(a)


def _attachment(db: Session, attachment_id: int, user: User, co: Company) -> tuple[Attachment, Entity]:
    a = db.get(Attachment, attachment_id)
    if not a or a.company_id != co.company_id:
        raise HTTPException(404, "Attachment not found")
    return a, _entity(a.entity, user)


@router.get("/attachment/{attachment_id}")
def download_attachment(attachment_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user),
                        co: Company = Depends(get_company)):
    a, _ = _attachment(db, attachment_id, user, co)
    path = os.path.join(get_settings().attachments_dir, a.stored_as)
    if not os.path.isfile(path):
        raise HTTPException(410, "File is missing on the server")
    return FileResponse(path, media_type=a.content_type, filename=a.filename)


@router.delete("/attachment/{attachment_id}", status_code=204)
def delete_attachment(attachment_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user),
                      co: Company = Depends(get_company)):
    a, ent = _attachment(db, attachment_id, user, co)
    if f"{ent.module}.delete" not in (user.role.permissions or []) and a.uploaded_by != user.user_id:
        raise HTTPException(403, f"Missing permission: {ent.module}.delete")
    path = os.path.join(get_settings().attachments_dir, a.stored_as)
    if os.path.isfile(path):
        os.remove(path)
    audit.log(db, user.user_id, "attachment", a.attachment_id, "delete", {"file": a.filename}, None)
    db.delete(a)
    db.commit()
