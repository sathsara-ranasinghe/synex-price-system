"""Bulk import from Excel / CSV for every registry entity.

Template columns come from the registry. Transactions with lines use a "Doc key" column: rows with the same key
become one document (header values from its first row, one line per row). References (customer, item, account …)
are written by name and resolved against the company's QuickBooks lists.
"""

import csv
import io
from dataclasses import dataclass
from datetime import date, datetime

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from ..database import get_db
from ..deps import get_company, get_current_user
from ..models import Company, QBRecord, User
from ..qb import builder
from ..qb.registry import BY_KEY, REF_GROUPS, Entity, F, LineType
from .portal import _describe, _queue

router = APIRouter(prefix="/api/files", tags=["Bulk import"])

MAX_ROWS = 5000
ADDRESS_COLS = [("Addr1", "Address line 1"), ("Addr2", "Address line 2"), ("City", "City"), ("State", "State"),
                ("PostalCode", "Postal code"), ("Country", "Country")]
TYPE_HELP = {"str": "Text", "text": "Text", "int": "Whole number", "decimal": "Number", "money": "Amount (e.g. 1250.00)",
             "date": "Date (YYYY-MM-DD)", "bool": "Yes / No", "enum": "One of the listed values",
             "ref": "Name exactly as in QuickBooks", "txnref": "Document number (Ref no.)"}
EXAMPLES = {"str": "Sample", "text": "Sample text", "int": "30", "decimal": "2", "money": "1500.00", "date": "2026-10-08",
            "bool": "Yes", "txnref": "INV-1001"}


@dataclass
class Col:
    header: str
    kind: str  # key | field | address | linetype | line
    field: F | None = None
    part: str | None = None  # address part
    line_label: str | None = None  # for kind=line: the field label shared by line types


def _entity(key: str, user: User) -> Entity:
    ent = BY_KEY.get(key)
    if not ent:
        raise HTTPException(404, "Unknown entity")
    if not ent.can_add:
        raise HTTPException(422, f"{ent.plural} cannot be created through QuickBooks Web Connector")
    if f"{ent.module}.create" not in (user.role.permissions or []):
        raise HTTPException(403, f"Missing permission: {ent.module}.create")
    return ent


def hidden_features(company: Company) -> set[str]:
    """Features this company file does not use, from the QuickBooks preferences read on sync."""
    prefs = company.preferences or {}
    hide = set()
    if builder.get(prefs, "AccountingPreferences/IsUsingClassTracking") == "false":
        hide.add("class")
    if builder.get(prefs, "MultiCurrencyPreferences/IsMultiCurrencyOn") == "false" \
            or "currency" in (company.disabled_entities or []):
        hide.add("currency")
    return hide


def _hidden(f: F, hide: set[str]) -> bool:
    return ("class" in hide and f.ref == "class") or \
        ("currency" in hide and (f.ref == "currency" or f.path == "ExchangeRate"))


def columns(ent: Entity, hide: set[str] | frozenset = frozenset()) -> list[Col]:
    cols: list[Col] = []
    if ent.lines:
        cols.append(Col("Doc key", "key"))
    used = set()
    for f in ent.fields:
        if not f.add or _hidden(f, hide):
            continue
        if f.type == "address":
            for part, label in ADDRESS_COLS:
                cols.append(Col(f"{f.label} - {label}", "address", f, part))
            continue
        header = f.label if f.label not in used else f"{f.label} ({f.path})"
        used.add(header)
        cols.append(Col(header, "field", f))
    if len(ent.lines) > 1:
        cols.append(Col("Line type", "linetype"))
    seen = set()
    for lt in ent.lines:
        for f in lt.fields:
            if f.label not in seen and not _hidden(f, hide):
                seen.add(f.label)
                cols.append(Col(f"Line - {f.label}", "line", f, line_label=f.label))
    return cols


def _line_field(lt: LineType, label: str) -> F | None:
    return next((f for f in lt.fields if f.label == label), None)


# ---------------------------------------------------------------- template

DROPDOWN_ROWS = MAX_ROWS + 1  # data validation range on the first sheet
MAX_LIST_NAMES = 20000


def _names(db: Session, company_id: int, f: F) -> list[str]:
    targets = REF_GROUPS.get(f.ref, [f.ref])
    q = db.query(QBRecord.name).filter(QBRecord.company_id == company_id, QBRecord.entity.in_(targets),
                                       QBRecord.deleted.is_(False), QBRecord.name.isnot(None))
    if BY_KEY[targets[0]].kind == "list":
        q = q.filter(QBRecord.is_active.is_(True))
    return sorted({n for (n,) in q.limit(MAX_LIST_NAMES)}, key=str.lower)


@router.get("/template/{entity}")
def template(entity: str, format: str = Query("xlsx", pattern="^(xlsx|csv)$"), db: Session = Depends(get_db),
             user: User = Depends(get_current_user), co: Company = Depends(get_company)):
    ent = _entity(entity, user)
    cols = columns(ent, hidden_features(co))
    example = []
    for c in cols:
        if c.kind == "key":
            example.append("DOC-1")
        elif c.kind == "linetype":
            example.append(ent.lines[-1].key)
        elif c.kind == "address":
            example.append({"Addr1": "No. 12, Main Street", "City": "Colombo", "Country": "Sri Lanka"}.get(c.part, ""))
        elif c.field.type == "enum":
            example.append(c.field.options[0])
        elif c.field.type == "ref":
            example.append(f"<{c.field.label} name>" if c.field.required else "")
        else:
            example.append(EXAMPLES.get(c.field.type, "") if (c.field.required or c.kind == "line") else "")
    fname = f"{ent.key}_import_template"
    if format == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow([c.header for c in cols])
        buf.seek(0)
        return StreamingResponse(io.BytesIO(("﻿" + buf.getvalue()).encode("utf-8")), media_type="text/csv",
                                 headers={"Content-Disposition": f"attachment; filename={fname}.csv"})

    wb = Workbook()
    ws = wb.active
    ws.title = ent.plural[:31]
    ws.append([c.header for c in cols])
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i)
        required = (c.kind == "key") or (c.field is not None and c.field.required and c.kind != "address")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0A5C80" if required else "5B7F95")
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(40, len(c.header) + 4))
    ws.freeze_panes = "A2"

    # dropdowns: QuickBooks names of THIS company for references, fixed values for enums / yes-no / line type
    lists = wb.create_sheet("Lists")
    list_col = 0
    for i, c in enumerate(cols, start=1):
        letter = get_column_letter(i)
        target = f"{letter}2:{letter}{DROPDOWN_ROWS}"
        options: list[str] = []
        warn = False
        if c.kind == "linetype":
            options = [lt.key for lt in ent.lines]
        elif c.field is not None and c.kind != "address":
            if c.field.type == "enum":
                options = list(c.field.options)
            elif c.field.type == "bool":
                options = ["Yes", "No"]
            elif c.field.type in ("ref", "txnref"):
                options, warn = _names(db, co.company_id, c.field), True
        if not options:
            continue
        list_col += 1
        lcol = get_column_letter(list_col)
        lists.cell(row=1, column=list_col, value=c.header)
        for r, name in enumerate(options, start=2):
            lists.cell(row=r, column=list_col, value=name)
        dv = DataValidation(type="list", formula1=f"={quote_sheetname('Lists')}!${lcol}$2:${lcol}${len(options) + 1}",
                            allow_blank=True, showErrorMessage=True,
                            errorStyle="warning" if warn else "stop",
                            errorTitle="Not in the list",
                            error=("This name is not in QuickBooks yet. Keep it only if it will exist before you import."
                                   if warn else "Pick a value from the list."))
        ws.add_data_validation(dv)
        dv.add(target)
    lists.sheet_state = "hidden"

    ex = wb.create_sheet("Example")
    ex.append([c.header for c in cols])
    ex.append(example)
    if ent.lines:  # a second line of the same document
        ex.append(["DOC-1" if c.kind == "key" else (example[i] if c.kind in ("line", "linetype") else "")
                   for i, c in enumerate(cols)])
    for i in range(1, len(cols) + 1):
        ex.cell(row=1, column=i).font = Font(bold=True)
        ex.column_dimensions[get_column_letter(i)].width = max(14, min(40, len(cols[i - 1].header) + 4))

    hp = wb.create_sheet("Help")
    hp.append(["Column", "Required", "Type", "Allowed values / notes"])
    for c in hp[1]:
        c.font = Font(bold=True)
    for c in cols:
        if c.kind == "key":
            hp.append([c.header, "Yes", "Text", "Rows with the same key become ONE document; each row is one line."])
        elif c.kind == "linetype":
            hp.append([c.header, "Yes", "Text", "One of: " + ", ".join(lt.key for lt in ent.lines)])
        else:
            f = c.field
            notes = ", ".join(f.options) if f.type == "enum" else ""
            if f.type in ("ref", "txnref"):
                target = REF_GROUPS.get(f.ref, [f.ref])
                notes = "Must already exist in QuickBooks: " + ", ".join(BY_KEY[t].plural for t in target)
            if f.max:
                notes = (notes + "; " if notes else "") + f"max {f.max} characters"
            if f.min is not None:
                notes = (notes + "; " if notes else "") + f"not below {f.min:g}"
            hp.append([c.header, "Yes" if f.required and c.kind != "address" else "", TYPE_HELP.get(f.type, f.type), notes])
    hp.append([])
    hp.append(["Fill the first sheet (one row per record" + ("/line" if ent.lines else "") +
               "). Leave optional columns empty. Names must match QuickBooks exactly (case does not matter)."])
    for col, width in zip("ABCD", (34, 10, 26, 90)):
        hp.column_dimensions[col].width = width
    for row in hp.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename={fname}.xlsx"})


# ---------------------------------------------------------------- parsing

def _read(file: UploadFile, content: bytes) -> list[dict]:
    name = (file.filename or "").lower()
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
    elif name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    else:
        raise HTTPException(400, "Upload an .xlsx or .csv file")
    if not rows:
        raise HTTPException(422, "The file is empty")
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = []
    for n, r in enumerate(rows[1:], start=2):
        if not any(v not in (None, "") and str(v).strip() for v in r):
            continue
        out.append({"__row": n, **{header[i]: r[i] for i in range(min(len(header), len(r))) if header[i]}})
    if len(out) > MAX_ROWS:
        raise HTTPException(413, f"At most {MAX_ROWS} rows per import")
    return out


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _convert(f: F, raw, errors: list[str], resolve) -> object:
    if isinstance(raw, (datetime, date)) and f.type == "date":
        return (raw.date() if isinstance(raw, datetime) else raw).isoformat()
    s = _cell(raw)
    if s == "":
        return None
    if f.type == "bool":
        v = s.lower()
        if v in ("yes", "y", "true", "1"):
            return True
        if v in ("no", "n", "false", "0"):
            return False
        errors.append(f"{f.label}: use Yes or No")
        return None
    if f.type == "date":
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
            try:
                return datetime.strptime(s[:10], fmt).date().isoformat()
            except ValueError:
                pass
        errors.append(f"{f.label}: '{s}' is not a date (use YYYY-MM-DD)")
        return None
    if f.type in ("money", "decimal", "int"):
        return s.replace(",", "")
    if f.type == "enum":
        match = next((o for o in f.options if o.lower() == s.replace(" ", "").lower()), None)
        if not match:
            errors.append(f"{f.label}: '{s}' must be one of {', '.join(f.options)}")
        return match
    if f.type in ("ref", "txnref"):
        found = resolve(f, s)
        if not found:
            errors.append(f"{f.label}: '{s}' not found in QuickBooks")
            return None
        return found["ListID"] if f.type == "txnref" else found
    return s


def _resolver(db: Session, company_id: int):
    cache: dict[tuple[str, str], dict | None] = {}

    def resolve(f: F, name: str):
        key = (f.ref, name.lower())
        if key not in cache:
            targets = REF_GROUPS.get(f.ref, [f.ref])
            rec = (db.query(QBRecord).filter(QBRecord.company_id == company_id, QBRecord.entity.in_(targets),
                                             QBRecord.deleted.is_(False))
                   .filter(or_(func.lower(QBRecord.name) == name.lower(), func.lower(QBRecord.qb_id) == name.lower()))
                   .first())
            if rec is None:  # sub-items / sub-accounts: match the short Name too ("Parent:Child" in FullName)
                for r in db.query(QBRecord).filter(QBRecord.company_id == company_id, QBRecord.entity.in_(targets),
                                                   QBRecord.deleted.is_(False)).limit(20000):
                    if str((r.data or {}).get("Name", "")).lower() == name.lower():
                        rec = r
                        break
            cache[key] = {"ListID": rec.qb_id, "FullName": rec.name} if rec else None
        return cache[key]

    return resolve


def parse_documents(db: Session, ent: Entity, company_id: int, rows: list[dict]) -> list[dict]:
    cols = columns(ent)
    resolve = _resolver(db, company_id)
    known = {c.header for c in cols}
    missing_required = [c.header for c in cols if c.kind == "key" or (c.kind == "field" and c.field.required)]

    groups: dict[str, list[dict]] = {}
    for r in rows:
        key = _cell(r.get("Doc key")) if ent.lines else ""
        groups.setdefault(key or f"__row{r['__row']}", []).append(r)

    docs = []
    for key, grp in groups.items():
        first = grp[0]
        errors: list[str] = []
        unknown = [h for h in first if h != "__row" and h not in known]
        if unknown and len(docs) == 0:
            errors.append("Unknown column(s): " + ", ".join(unknown[:5]) + " - use the template")
        values: dict = {}
        for c in cols:
            if c.kind == "field":
                v = _convert(c.field, first.get(c.header), errors, resolve)
                if v is not None:
                    values[c.field.path] = v
            elif c.kind == "address":
                v = _cell(first.get(c.header))
                if v:
                    values.setdefault(c.field.path, {})[c.part] = v
        lines = []
        for r in grp:
            if not ent.lines:
                break
            lt_key = _cell(r.get("Line type")) if len(ent.lines) > 1 else ent.lines[0].key
            lt = next((x for x in ent.lines if x.key.lower() == lt_key.lower()), None)
            raw_line = {c.line_label: r.get(c.header) for c in cols if c.kind == "line" and _cell(r.get(c.header))}
            if not raw_line:
                continue
            if lt is None:
                errors.append(f"Row {r['__row']}: Line type must be one of {', '.join(x.key for x in ent.lines)}")
                continue
            lv = {}
            for label, raw in raw_line.items():
                f = _line_field(lt, label)
                if f is None:
                    errors.append(f"Row {r['__row']}: '{label}' is not used on {lt.label.lower()} lines")
                    continue
                v = _convert(f, raw, errors, resolve)
                if v is not None:
                    lv[f.path] = v
            lines.append({"type": lt.key, "values": lv})
        if not errors:
            try:
                builder.build_add(ent, values, lines or None, "check")
            except builder.BuildError as e:
                errors.append(str(e))
        docs.append({"row": first["__row"], "key": None if key.startswith("__row") else key,
                     "summary": _describe(ent, values) or (key if not key.startswith("__row") else f"Row {first['__row']}"),
                     "lines": len(lines), "values": values, "line_items": lines, "errors": errors})
    if not rows:
        raise HTTPException(422, "No data rows found. Required columns: " + ", ".join(missing_required))
    return docs


# ---------------------------------------------------------------- preview / commit

@router.post("/import/{entity}/preview")
async def preview(entity: str, file: UploadFile = File(...), db: Session = Depends(get_db),
                  user: User = Depends(get_current_user), co: Company = Depends(get_company)):
    ent = _entity(entity, user)
    content = await file.read(15 * 1024 * 1024 + 1)
    if len(content) > 15 * 1024 * 1024:
        raise HTTPException(413, "File larger than 15 MB")
    docs = parse_documents(db, ent, co.company_id, _read(file, content))
    ok = sum(1 for d in docs if not d["errors"])
    return {"entity": ent.key, "documents": docs, "valid": ok, "invalid": len(docs) - ok,
            "direct": f"{ent.module}.direct" in (user.role.permissions or [])}


class CommitDoc(BaseModel):
    values: dict
    line_items: list[dict] = []


class CommitIn(BaseModel):
    documents: list[CommitDoc]


@router.post("/import/{entity}/commit")
def commit(entity: str, body: CommitIn, db: Session = Depends(get_db), user: User = Depends(get_current_user),
           co: Company = Depends(get_company)):
    ent = _entity(entity, user)
    if len(body.documents) > MAX_ROWS:
        raise HTTPException(413, f"At most {MAX_ROWS} documents per import")
    queued = approved = 0
    errors = []
    for i, d in enumerate(body.documents, start=1):
        lines = d.line_items or None
        try:  # validate again - never trust the browser copy
            builder.build_add(ent, d.values, lines, "check")
        except builder.BuildError as e:
            errors.append(f"Document {i}: {e}")
            continue
        w = _queue(db, user, co, ent, "qb_add", f"Import: new {ent.label.lower()} {_describe(ent, d.values)}",
                   {"values": d.values, "lines": lines})
        if w.status == "approved":
            approved += 1
        else:
            queued += 1
    return {"sent_to_quickbooks": approved, "waiting_for_approval": queued, "skipped": len(errors), "errors": errors[:50]}
