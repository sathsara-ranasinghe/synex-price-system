"""Bulk import: templates for every entity, CSV/Excel preview with name resolution and validation, commit."""

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.database import SessionLocal
from app.main import app
from app.models import QBRecord
from app.qb.registry import ENTITIES
from app.routers.imports import columns

client = TestClient(app)
client.__enter__()
H = {"Authorization": "Bearer " + client.post("/api/auth/login", data={"username": "admin", "password": "Admin@123"}
                                              ).json()["access_token"], "X-Company-Id": "1"}



def seed(entity: str, qb_id: str, name: str):
    with SessionLocal() as db:
        if not db.query(QBRecord).filter_by(company_id=1, entity=entity, qb_id=qb_id).first():
            db.add(QBRecord(company_id=1, entity=entity, qb_id=qb_id, name=name, data={"Name": name.split(":")[-1]}))
            db.commit()


@pytest.fixture(autouse=True, scope="module")
def lists():
    seed("customer", "IMP-C1", "Import Customer")
    seed("item_service", "IMP-I1", "Consulting:Design Work")
    seed("account", "IMP-A1", "Rent Expense")
    seed("vendor", "IMP-V1", "Import Vendor")


@pytest.mark.parametrize("key", [e.key for e in ENTITIES if e.can_add])
def test_template_for_every_entity(key):
    xl = client.get(f"/api/files/template/{key}", headers=H)
    assert xl.status_code == 200
    wb = load_workbook(io.BytesIO(xl.content))
    assert wb.sheetnames[1:] == ["Lists", "Example", "Help"] and wb["Lists"].sheet_state == "hidden"
    header = [c.value for c in wb.worksheets[0][1]]
    from app.models import Company
    from app.routers.imports import hidden_features
    with SessionLocal() as db:
        hide = hidden_features(db.get(Company, 1))
    assert header == [c.header for c in columns(next(e for e in ENTITIES if e.key == key), hide)]
    csv_ = client.get(f"/api/files/template/{key}", headers=H, params={"format": "csv"})
    assert csv_.status_code == 200 and csv_.content.decode("utf-8-sig").splitlines()[0].split(",")[0] == header[0]


def test_csv_customers_preview_and_commit():
    data = ("Customer name,Company,Phone,Email,Credit limit,Bill to - Address line 1,Bill to - City\n"
            "Alpha Hotels,Alpha (Pvt) Ltd,0112 000 000,a@alpha.lk,50000,1 Galle Rd,Colombo\n"
            ",No Name Ltd,,,,,\n"
            "Beta Stores,,,,-10,,\n")
    r = client.post("/api/files/import/customer/preview", headers=H, files={"file": ("c.csv", data, "text/csv")}).json()
    assert r["valid"] == 1 and r["invalid"] == 2
    good = r["documents"][0]
    assert good["values"]["BillAddress"] == {"Addr1": "1 Galle Rd", "City": "Colombo"} and not good["errors"]
    assert any("required" in e for e in r["documents"][1]["errors"])
    assert any("less than 0" in e for e in r["documents"][2]["errors"])
    res = client.post("/api/files/import/customer/commit", headers=H,
                      json={"documents": [d for d in r["documents"] if not d["errors"]]}).json()
    assert res["sent_to_quickbooks"] + res["waiting_for_approval"] == 1 and res["skipped"] == 0


def test_excel_invoices_grouped_by_doc_key():
    wb = Workbook()
    ws = wb.active
    ws.append(["Doc key", "Customer:Job", "Date", "Invoice no.", "Memo", "Line - Item", "Line - Description",
               "Line - Qty", "Line - Rate"])
    ws.append(["A", "import customer", "2026-10-08", "IMP-1", "Two lines", "Design Work", "Concept", 2, 1500])
    ws.append(["A", None, None, None, None, "Consulting:Design Work", "Drawings", 1, 800])
    ws.append(["B", "Unknown Customer", "08/10/2026", "IMP-2", None, None, "Note only line", None, None])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/files/import/invoice/preview", headers=H,
                    files={"file": ("inv.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
                    ).json()
    a, b = r["documents"]
    assert a["key"] == "A" and a["lines"] == 2 and not a["errors"], a
    assert a["values"]["CustomerRef"] == {"ListID": "IMP-C1", "FullName": "Import Customer"}
    assert [ln["values"]["ItemRef"]["ListID"] for ln in a["line_items"]] == ["IMP-I1", "IMP-I1"]
    assert a["values"]["TxnDate"] == "2026-10-08"
    assert b["values"]["TxnDate"] == "2026-10-08"  # dd/mm/yyyy understood
    assert any("Unknown Customer" in e for e in b["errors"])


def test_bill_with_expense_and_item_lines():
    data = ("Doc key,Vendor,Ref no.,Line type,Line - Account,Line - Amount,Line - Item,Line - Qty,Line - Cost\n"
            "X,Import Vendor,B-1,expense,Rent Expense,25000,,,\n"
            "X,,,item,,,Design Work,3,100\n"
            "X,,,wrong,,,,,5\n")
    r = client.post("/api/files/import/bill/preview", headers=H, files={"file": ("b.csv", data, "text/csv")}).json()
    doc = r["documents"][0]
    assert any("Line type must be one of" in e for e in doc["errors"])
    fixed = data.replace("X,,,wrong,,,,,5\n", "")
    doc = client.post("/api/files/import/bill/preview", headers=H, files={"file": ("b.csv", fixed, "text/csv")}).json()["documents"][0]
    assert not doc["errors"] and [ln["type"] for ln in doc["line_items"]] == ["expense", "item"]


def test_import_rules():
    assert client.post("/api/files/import/customer/preview", headers=H,
                       files={"file": ("c.txt", "x", "text/plain")}).status_code == 400
    assert client.get("/api/files/template/employee", headers=H).status_code == 200
    assert client.get("/api/files/template/no_such", headers=H).status_code == 404
    bad = client.post("/api/files/import/customer/commit", headers=H, json={"documents": [{"values": {}}]}).json()
    assert bad["skipped"] == 1 and "required" in bad["errors"][0]
    # a viewer cannot import
    client.post("/api/roles", headers=H, json={"role_name": "imp_viewer", "permissions": ["sales.view"]})
    client.post("/api/users", headers=H, json={"username": "impviewer", "email": "i@x.lk", "password": "Passw0rd!",
                                               "role_name": "imp_viewer", "company_ids": [1]})
    hv = {"Authorization": "Bearer " + client.post("/api/auth/login", data={"username": "impviewer", "password": "Passw0rd!"}
                                                   ).json()["access_token"]}
    assert client.get("/api/files/template/customer", headers=hv).status_code == 403


def test_template_follows_company_file():
    from app.models import Company
    with SessionLocal() as db:
        co = db.get(Company, 1)
        old = co.preferences
        co.preferences = {"AccountingPreferences": {"IsUsingClassTracking": "false"},
                          "MultiCurrencyPreferences": {"IsMultiCurrencyOn": "false"}}
        db.commit()
    try:
        wb = load_workbook(io.BytesIO(client.get("/api/files/template/invoice", headers=H).content))
        header = [c.value for c in wb.worksheets[0][1]]
        assert "Class" not in header and "Line - Class" not in header and not any("Exchange rate" in h for h in header)
        assert "Customer:Job" in header
        # dropdowns: customers of this company in the hidden Lists sheet, validation on the Customer column
        lists = wb["Lists"]
        cols = {lists.cell(row=1, column=i).value: i for i in range(1, lists.max_column + 1)}
        customers = [lists.cell(row=r, column=cols["Customer:Job"]).value for r in range(2, lists.max_row + 1)]
        assert "Import Customer" in customers
        ws = wb.worksheets[0]
        cust_letter = chr(ord("A") + header.index("Customer:Job"))
        assert any(str(dv.sqref).startswith(f"{cust_letter}2") for dv in ws.data_validations.dataValidation)
        # previously exported columns are still accepted on import
        data = "Customer name,Currency (multi-currency files)\nGamma Ltd,\n"
        r = client.post("/api/files/import/customer/preview", headers=H, files={"file": ("c.csv", data, "text/csv")}).json()
        assert r["valid"] == 1
    finally:
        with SessionLocal() as db:
            db.get(Company, 1).preferences = old
            db.commit()


def test_preferences_are_synced():
    from app.models import Company
    from app.qb import store
    from lxml import etree
    ret = etree.fromstring("<PreferencesRet><AccountingPreferences><IsUsingClassTracking>true</IsUsingClassTracking>"
                           "</AccountingPreferences></PreferencesRet>")
    with SessionLocal() as db:
        old = db.get(Company, 1).preferences
        store.process(db, [ret], 1)
        db.commit()
        assert db.get(Company, 1).preferences["AccountingPreferences"]["IsUsingClassTracking"] == "true"
        db.get(Company, 1).preferences = old
        db.commit()
